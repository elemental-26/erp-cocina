import json
import math
import os
import re
import sys
import unicodedata
from collections import defaultdict

import openpyxl


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = r"C:\Users\Usuario\Downloads\Estandarizacion_ cocina _v_01.xlsx"
OUT = os.path.join(ROOT, "src", "data", "estandarizacionCocinaData.js")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.replace("\n", " ").split()).strip()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and math.isnan(value):
            return ""
        return value
    return str(value).strip()


def text(value):
    value = clean(value)
    return "" if value is None else str(value).strip()


def number(value, default=0):
    value = clean(value)
    if value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def slug(value):
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-zA-Z0-9]+", "-", value.lower()).strip("-")
    return value or "item"


def normalize_name(value):
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", value.lower()).strip()


def pct(value):
    raw = number(value)
    if raw > 1:
        raw = raw / 100
    return round(raw, 4)


def next_code(prefix, used):
    prefix = re.sub(r"[^A-Z]", "", text(prefix).upper())[:3] or "GEN"
    index = 1
    while f"{prefix}{index:03d}" in used:
        index += 1
    code = f"{prefix}{index:03d}"
    used.add(code)
    return code


def get_sheet(wb, name):
    for sheet in wb.worksheets:
        if sheet.title.strip().upper() == name.strip().upper():
            return sheet
    return None


def export(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    familias = []
    prefix_to_family = {}
    ws = get_sheet(wb, "FAMILIAS")
    if ws:
        for row in ws.iter_rows(min_row=2, values_only=True):
            nombre, prefijo = text(row[0]), text(row[1])
            if not nombre or not prefijo:
                continue
            item = {"id": slug(nombre), "nombre": nombre.title(), "prefijo": prefijo.upper()}
            familias.append(item)
            prefix_to_family[item["prefijo"]] = item["nombre"]

    insumos = []
    seen_names = set()
    used_codes = set()
    ws = get_sheet(wb, "INSUMOS")
    if ws:
        for row in ws.iter_rows(min_row=7, values_only=True):
            raw_id, nombre = text(row[0]).upper(), text(row[1])
            if not nombre:
                continue
            norm = normalize_name(nombre)
            if norm in seen_names:
                continue
            seen_names.add(norm)
            prefix = re.match(r"^[A-Z]+", raw_id or "")
            prefix = prefix.group(0) if prefix else "GEN"
            if raw_id and raw_id not in used_codes:
                item_id = raw_id
                used_codes.add(item_id)
            else:
                item_id = next_code(prefix, used_codes)
            presentacion = number(row[6])
            valor_compra = number(row[7])
            merma = pct(row[4])
            costo_unitario = round(valor_compra / presentacion, 4) if presentacion else number(row[3])
            costo_real = round(costo_unitario / (1 - merma), 4) if costo_unitario and merma < 1 else costo_unitario
            insumos.append({
                "id": item_id,
                "originalId": raw_id,
                "nombre": nombre,
                "familia": prefix_to_family.get(prefix, "General"),
                "unidad": text(row[2]) or "und",
                "costoUnitario": costo_unitario,
                "mermaPct": merma,
                "costoReal": costo_real,
                "presentacion": presentacion,
                "valorCompra": valor_compra,
                "activo": text(row[8]).lower() != "no",
            })

    insumo_by_name = {normalize_name(item["nombre"]): item for item in insumos}

    calidad = {}
    ws = get_sheet(wb, "CALIDAD")
    if ws:
        for row in ws.iter_rows(min_row=2, values_only=True):
            rid = text(row[0]).upper()
            if not rid:
                continue
            calidad[rid] = {
                "color": text(row[1]),
                "textura": text(row[2]),
                "sabor": text(row[3]),
                "aroma": text(row[4]),
                "temperaturaCoccion": text(row[5]),
                "temperaturaConservacion": text(row[6]),
                "presentacion": text(row[7]),
                "rangos": text(row[8]),
                "tolerancias": text(row[9]),
                "criteriosRechazo": text(row[10]),
                "codigoFicha": text(row[11]),
                "complejidad": text(row[12]),
                "servicio": text(row[13]),
            }

    procedimientos = defaultdict(lambda: {"pasos": [], "observacionesGenerales": ""})
    ws = get_sheet(wb, "PROCEDIMIENTO")
    if ws:
        for row in ws.iter_rows(min_row=2, values_only=True):
            rid = text(row[0]).upper()
            if not rid:
                continue
            if text(row[7]) and not procedimientos[rid]["observacionesGenerales"]:
                procedimientos[rid]["observacionesGenerales"] = text(row[7])
            if text(row[2]):
                procedimientos[rid]["pasos"].append({
                    "paso": text(row[1]),
                    "descripcion": text(row[2]),
                    "tiempoMin": number(row[3], ""),
                    "temperatura": text(row[4]),
                    "equipo": text(row[5]),
                    "observaciones": text(row[6]),
                })

    recetas_map = {}
    ws = get_sheet(wb, "RECETAS")
    if ws:
        for row in ws.iter_rows(min_row=2, values_only=True):
            rid, nombre = text(row[0]).upper(), text(row[1])
            if not rid and not nombre:
                continue
            if not rid:
                rid = f"REC-{slug(nombre).upper()}"
            receta = recetas_map.setdefault(rid, {
                "id": rid,
                "nombre": nombre or rid,
                "version": "1",
                "creadoPor": "",
                "rendimiento": "",
                "pax": int(number(row[7], 0)) if number(row[7], 0) else "",
                "pesoPorPax": "",
                "foto": "",
                "ingredientes": [],
                "procedimiento": [],
                "calidad": {},
                "observacionesGenerales": "",
            })
            if nombre and receta["nombre"] == rid:
                receta["nombre"] = nombre
            if not receta["pax"] and number(row[7], 0):
                receta["pax"] = int(number(row[7], 0))
            insumo_nombre = text(row[2])
            if insumo_nombre:
                found = insumo_by_name.get(normalize_name(insumo_nombre), {})
                receta["ingredientes"].append({
                    "id": f"ing-{slug(rid)}-{len(receta['ingredientes']) + 1}",
                    "insumoId": found.get("id", text(row[3]).upper()),
                    "insumo": insumo_nombre,
                    "cantidad": number(row[4], ""),
                    "unidad": text(row[5]),
                })

    produccion = get_sheet(wb, "PRODUCCION")
    if produccion:
        for row in produccion.iter_rows(min_row=2, values_only=True):
            rid = text(row[0]).upper()
            if rid in recetas_map and number(row[1], 0):
                recetas_map[rid]["pax"] = int(number(row[1], 0))

    for rid, receta in recetas_map.items():
        receta["procedimiento"] = procedimientos[rid]["pasos"]
        receta["observacionesGenerales"] = procedimientos[rid]["observacionesGenerales"]
        receta["calidad"] = calidad.get(rid, {})

    preparaciones = []
    seen_prep = set()
    ws = get_sheet(wb, "BD-PREPARACIONES")
    if ws:
        for row in ws.iter_rows(min_row=9, values_only=True):
            codigo, categoria, subcategoria, nombre = text(row[0]).upper(), text(row[1]), text(row[2]), text(row[3])
            if not nombre:
                continue
            key = codigo or normalize_name(nombre)
            if key in seen_prep:
                continue
            seen_prep.add(key)
            preparaciones.append({
                "id": codigo or f"PREP-{len(preparaciones) + 1:03d}",
                "codigo": codigo,
                "categoria": categoria,
                "subcategoria": subcategoria,
                "nombre": nombre,
                "estado": text(row[4]) or "Activo",
                "observaciones": text(row[5]),
            })

    data = {
        "source": os.path.basename(path),
        "familias": familias,
        "insumos": insumos,
        "recetas": sorted(recetas_map.values(), key=lambda item: item["id"]),
        "preparaciones": preparaciones,
        "mermas": [],
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as handle:
        handle.write("// Datos iniciales extraidos de Estandarizacion_ cocina _v_01.xlsx.\n")
        handle.write("// El modulo conserva los cambios del usuario en localStorage y no reimporta encima.\n")
        handle.write("export const DEFAULT_STD_DATA = ")
        json.dump(data, handle, ensure_ascii=False, indent=2)
        handle.write(";\n")
    return data


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    result = export(src)
    print(json.dumps({
        "familias": len(result["familias"]),
        "insumos": len(result["insumos"]),
        "recetas": len(result["recetas"]),
        "preparaciones": len(result["preparaciones"]),
        "salida": OUT,
    }, ensure_ascii=True))
